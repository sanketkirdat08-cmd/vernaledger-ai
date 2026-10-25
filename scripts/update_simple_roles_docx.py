import docx

def update_docx_roles(filename):
    doc = docx.Document(filename)
    
    # Replacement map
    rep_map = {
        "Sanket Kirdat (Lead Developer)": "Sanket Kirdat (Team Leader & AI Developer)",
        "Sakshi Bhagat (Developer)": "Sakshi Bhagat (Frontend & UI Developer)",
        "Vaishnavi Dhavale (Developer)": "Vaishnavi Dhavale (Voice Billing & AI Developer)",
        "Rushikesh Mulik (Developer)": "Rushikesh Mulik (Backend & Database Developer)"
    }

    # Update in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for old_val, new_val in rep_map.items():
                    if old_val in cell.text:
                        cell.text = cell.text.replace(old_val, new_val)
                        print(f"Updated in {filename} table: {new_val}")

    # Update in paragraphs
    for p in doc.paragraphs:
        for old_val, new_val in rep_map.items():
            if old_val in p.text:
                p.text = p.text.replace(old_val, new_val)
                print(f"Updated in {filename} paragraph: {new_val}")

    doc.save(filename)
    print(f"Successfully saved {filename}!")

update_docx_roles("VernaLedger_AI_Final_Year_Project_Report.docx")
update_docx_roles("VernaLedger_AI_Project_Synopsis.docx")

