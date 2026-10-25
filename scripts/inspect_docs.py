import sys
sys.stdout.reconfigure(encoding='utf-8')

import docx, pptx

print("=== 1. VernaLedger_AI_Final_Year_Project_Report.docx ===")
doc = docx.Document("VernaLedger_AI_Final_Year_Project_Report.docx")
for i, p in enumerate(doc.paragraphs[:50]):
    text = p.text.strip()
    if any(k in text for k in ["Sanket", "Sakshi", "Vaishnavi", "Rushikesh", "Roll", "PRN", "Submitted", "Students"]):
        print(f"P{i}: {text}")

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        row_text = " | ".join(c.text.strip() for c in row.cells)
        if any(k in row_text for k in ["Sanket", "Sakshi", "Vaishnavi", "Rushikesh", "Roll", "PRN", "Name", "Student"]):
            print(f"Table {t_idx} Row {r_idx}: {row_text}")

print("\n=== 2. VernaLedger_AI_Project_Synopsis.docx ===")
syn = docx.Document("VernaLedger_AI_Project_Synopsis.docx")
for i, p in enumerate(syn.paragraphs[:50]):
    text = p.text.strip()
    if any(k in text for k in ["Sanket", "Sakshi", "Vaishnavi", "Rushikesh", "Roll", "PRN", "Submitted", "Students"]):
        print(f"P{i}: {text}")

for t_idx, table in enumerate(syn.tables):
    for r_idx, row in enumerate(table.rows):
        row_text = " | ".join(c.text.strip() for c in row.cells)
        if any(k in row_text for k in ["Sanket", "Sakshi", "Vaishnavi", "Rushikesh", "Roll", "PRN", "Name", "Student"]):
            print(f"Table {t_idx} Row {r_idx}: {row_text}")

print("\n=== 3. VERNALEDGER_AI_MERCHANT_PRO_PRESENTATION.pptx ===")
prs = pptx.Presentation("VERNALEDGER_AI_MERCHANT_PRO_PRESENTATION.pptx")
for s_idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if any(k in text for k in ["Sanket", "Sakshi", "Vaishnavi", "Rushikesh", "Roll", "PRN"]):
                    print(f"Slide {s_idx + 1}: {text}")
