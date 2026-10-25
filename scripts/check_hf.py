import docx

for fname in ['VernaLedger_AI_Final_Year_Project_Report.docx', 'VernaLedger_AI_Project_Synopsis.docx']:
    d = docx.Document(fname)
    for s_idx, sec in enumerate(d.sections):
        h_text = sec.header.paragraphs[0].text if sec.header.paragraphs else ''
        f_text = sec.footer.paragraphs[0].text if sec.footer.paragraphs else ''
        if h_text or f_text:
            print(f'{fname} sec[{s_idx}] header: "{h_text}", footer: "{f_text}"')

