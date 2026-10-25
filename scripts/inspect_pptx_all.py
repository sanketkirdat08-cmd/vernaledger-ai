import sys
sys.stdout.reconfigure(encoding='utf-8')
import pptx

prs = pptx.Presentation("VERNALEDGER_AI_MERCHANT_PRO_PRESENTATION.pptx")
print(f"Total slides in PPTX: {len(prs.slides)}")

for s_idx, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if any(w in t for w in ["Kirdat", "Bhagat", "Dhawale", "Dhavale", "Mulik", "Roll", "PRN"]):
                    if "Dr. S.S. Mulik" not in t:
                        print(f"Slide {s_idx}: {t}")
        if shape.has_table:
            for row in shape.table.rows:
                row_t = " | ".join(c.text.strip() for c in row.cells)
                if any(w in row_t for w in ["Kirdat", "Bhagat", "Dhawale", "Dhavale", "Mulik", "Roll", "PRN"]):
                    if "Dr. S.S. Mulik" not in row_t:
                        print(f"Slide {s_idx} Table: {row_t}")
