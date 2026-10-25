import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

def audit_docx(filename):
    print(f"\n=================== AUDITING {filename} ===================")
    doc = docx.Document(filename)
    
    # Check paragraphs
    for idx, p in enumerate(doc.paragraphs):
        t = p.text
        if any(w in t for w in ["Kirdat", "Bhagat", "Dhawale", "Dhavale", "Mulik", "Roll", "PRN", "Enrollment", "242133", "303"]):
            print(f"P[{idx}]: {t}")

    # Check tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_text = " | ".join(c.text.strip() for c in row.cells)
            if any(w in row_text for w in ["Kirdat", "Bhagat", "Dhawale", "Dhavale", "Mulik", "Roll", "PRN", "Enrollment", "242133", "303", "40"]):
                print(f"Table[{t_idx}] Row[{r_idx}]: {row_text}")

audit_docx("VernaLedger_AI_Final_Year_Project_Report.docx")
audit_docx("VernaLedger_AI_Project_Synopsis.docx")

