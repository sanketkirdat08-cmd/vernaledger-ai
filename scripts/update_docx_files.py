import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

def set_cell_text(cell, text, bold=False, font_size=11, font_name="Times New Roman", align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = font_name
    run.font.size = docx.shared.Pt(font_size)

def update_report_docx():
    doc = docx.Document("VernaLedger_AI_Final_Year_Project_Report.docx")
    
    # Students data
    students = [
        {"roll": "4060", "prn": "24065451263508", "name": "Sanket Kirdat (Lead Developer)"},
        {"roll": "4054", "prn": "24065451263502", "name": "Sakshi Bhagat (Developer)"},
        {"roll": "4056", "prn": "24065451263504", "name": "Vaishnavi Dhavale (Developer)"},
        {"roll": "4004", "prn": "23065451263005", "name": "Rushikesh Mulik (Developer)"}
    ]

    # Table 0: Title Page Candidate Details
    t0 = doc.tables[0]
    set_cell_text(t0.rows[0].cells[0], "Roll No.", bold=True, font_size=11)
    set_cell_text(t0.rows[0].cells[1], "Enrollment / PRN No.", bold=True, font_size=11)
    set_cell_text(t0.rows[0].cells[2], "Name of students", bold=True, font_size=11)

    for i, s in enumerate(students, start=1):
        set_cell_text(t0.rows[i].cells[0], s["roll"], bold=False, font_size=11)
        set_cell_text(t0.rows[i].cells[1], s["prn"], bold=False, font_size=11)
        set_cell_text(t0.rows[i].cells[2], s["name"], bold=True if "Lead" in s["name"] else False, font_size=11)

    # Table 1: Certificate Page Candidate Details
    t1 = doc.tables[1]
    set_cell_text(t1.rows[0].cells[0], "Roll No.", bold=True, font_size=11)
    set_cell_text(t1.rows[0].cells[1], "Enrollment / PRN No.", bold=True, font_size=11)
    set_cell_text(t1.rows[0].cells[2], "Name of students", bold=True, font_size=11)

    for i, s in enumerate(students, start=1):
        set_cell_text(t1.rows[i].cells[0], s["roll"], bold=False, font_size=11)
        set_cell_text(t1.rows[i].cells[1], s["prn"], bold=False, font_size=11)
        set_cell_text(t1.rows[i].cells[2], s["name"], bold=True if "Lead" in s["name"] else False, font_size=11)

    # Table 3: Signature Page Candidate Details
    t3 = doc.tables[3]
    for i, s in enumerate(students, start=1):
        label = f"{s['name']}\n(Roll No: {s['roll']} | PRN: {s['prn']})"
        set_cell_text(t3.rows[i].cells[1], label, bold=False, font_size=10.5)

    doc.save("VernaLedger_AI_Final_Year_Project_Report.docx")
    print("Updated VernaLedger_AI_Final_Year_Project_Report.docx successfully.")

def update_synopsis_docx():
    doc = docx.Document("VernaLedger_AI_Project_Synopsis.docx")
    
    students = [
        {"roll": "4060", "prn": "24065451263508", "name": "Sanket Kirdat (Lead Developer)"},
        {"roll": "4054", "prn": "24065451263502", "name": "Sakshi Bhagat (Developer)"},
        {"roll": "4056", "prn": "24065451263504", "name": "Vaishnavi Dhavale (Developer)"},
        {"roll": "4004", "prn": "23065451263005", "name": "Rushikesh Mulik (Developer)"}
    ]

    t0 = doc.tables[0]
    set_cell_text(t0.rows[0].cells[0], "Roll No.", bold=True, font_size=11)
    set_cell_text(t0.rows[0].cells[1], "Enrollment / PRN No.", bold=True, font_size=11)
    set_cell_text(t0.rows[0].cells[2], "Name of students", bold=True, font_size=11)

    for i, s in enumerate(students, start=1):
        set_cell_text(t0.rows[i].cells[0], s["roll"], bold=False, font_size=11)
        set_cell_text(t0.rows[i].cells[1], s["prn"], bold=False, font_size=11)
        set_cell_text(t0.rows[i].cells[2], s["name"], bold=True if "Lead" in s["name"] else False, font_size=11)

    doc.save("VernaLedger_AI_Project_Synopsis.docx")
    print("Updated VernaLedger_AI_Project_Synopsis.docx successfully.")

update_report_docx()
update_synopsis_docx()
